from django import forms
from django.shortcuts import render, redirect
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy, reverse
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from store_management.models import Store , Department
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from io import BytesIO

from django.views.generic import TemplateView
from django.db.models import Q, Count
from django.contrib import messages

import csv
from openpyxl import Workbook

from .forms import DepartmentForm, StoreForm
from retail_management_system.mixins import PaginationMixin

import logging

logger = logging.getLogger(__name__)

# Store Views



def department_edit(request, pk):
    department = get_object_or_404(Department, pk=pk)
    
    if request.method == 'POST':
        form = DepartmentForm(request.POST, instance=department)
        if form.is_valid():
            form.save()
            return redirect('store_management:department_detail', pk=department.pk)
    else:
        form = DepartmentForm(instance=department)
    
    return render(request, 'departments/edit.html', {
        'form': form,
        'department': department
    })

class StoreListView(PaginationMixin, ListView):
    model = Store
    context_object_name = 'stores'
    template_name = 'store_management/store_list.html'
    paginate_by = 10

    def get_queryset(self):
        queryset = super().get_queryset()
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) |
                Q(city__icontains=search) |
                Q(region__icontains=search)
            )
        region = self.request.GET.get('region')
        if region:
            queryset = queryset.filter(region=region)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['region_list'] = Store.objects.values_list('region', flat=True).distinct()
        context['search_query'] = self.request.GET.get('search', '')
        context['selected_region'] = self.request.GET.get('region', '')
        context['total_stores'] = Store.objects.count()
        return context

class StoreDetailView(DetailView):
    model = Store
    context_object_name = 'store'
    template_name = 'store_management/store_detail.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Ensure we're getting a valid queryset with PKs
        context['departments'] = self.object.departments.all().order_by('department_name')
        return context

class StoreCreateView(CreateView):
    form_class = StoreForm  # Use custom form instead of fields
    template_name = 'store_management/store_form.html'
    success_url = reverse_lazy('store_management:store_list')

    def form_valid(self, form):
        try:
            return super().form_valid(form)
        except Exception as e:
            # Add error to form for display
            form.add_error(None, f"Save error: {str(e)}")
            return self.form_invalid(form)

class StoreUpdateView(UpdateView):
    model = Store
    form_class = StoreForm
    template_name = 'store_management/store_form.html'
    success_url = reverse_lazy('store_management:store_list')



class StoreDeleteView(DeleteView):
    model = Store
    template_name = 'store_management/store_confirm_delete.html'
    success_url = reverse_lazy('store_management:store_list')

# Department Views
# store_management/views.py
# views.py
from django.shortcuts import get_object_or_404

class DepartmentCreateView(CreateView):
    model = Department
    form_class = DepartmentForm

    def form_valid(self, form):
        # Use id instead of store_id
        store = get_object_or_404(Store, id=self.kwargs['store_id'])
        form.instance.store = store
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Use id here too
        context['store'] = get_object_or_404(Store, id=self.kwargs['store_id'])
        return context

    def get_success_url(self):
        return reverse('store_management:department_list', kwargs={
            'store_id': self.kwargs['store_id']
        })

class DepartmentCreateStandaloneView(CreateView):
    model = Department
    form_class = DepartmentForm
    template_name = 'store_management/department_form.html'
    success_url = reverse_lazy('store_management:all_departments_list')

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        # Ensure we create a new instance
        kwargs['instance'] = None
        return kwargs

    def form_valid(self, form):
        # Check if the user has a store
        if not hasattr(self.request.user, 'store') or not self.request.user.store:
            form.add_error(None, "You cannot create a department without a store.")
            return self.form_invalid(form)
        # Set the store to the logged-in user's store
        form.instance.store = self.request.user.store
        return super().form_valid(form)
# views.py
class DepartmentUpdateView(UpdateView):
    model = Department
    form_class = DepartmentForm
    template_name = 'store_management/department_edit.html'
    context_object_name = 'department'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Add store to context
        context['store'] = self.object.store
        return context
    
    def get_success_url(self):
        return reverse('store_management:department_list', kwargs={
            'store_id': self.object.store.id
        })
class DepartmentDeleteView(DeleteView):
    model = Department
    template_name = 'store_management/department_confirm_delete.html'

    def get_success_url(self):
        # Use store_id from department's foreign key
        return reverse_lazy('store_management:department_list', kwargs={'store_id': self.object.store.id})


class DepartmentListView(PaginationMixin, ListView):
    model = Department
    context_object_name = 'departments'
    template_name = 'store_management/department_list.html'
    paginate_by = 10

    def get_queryset(self):
        store_id = self.kwargs['store_id']
        # Use store_id field directly
        return Department.objects.filter(store_id=store_id)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        store_id = self.kwargs['store_id']
        # Use id field instead of store_id
        context['store'] = get_object_or_404(Store, id=store_id)
        return context

class AllDepartmentsListView(PaginationMixin, ListView):
    model = Department
    context_object_name = 'departments'
    template_name = 'store_management/all_departments_list.html'
    paginate_by = 10

    def get_queryset(self):
        queryset = Department.objects.select_related('store').order_by('store__name', 'department_name')
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(department_name__icontains=search) |
                Q(description__icontains=search) |
                Q(store__name__icontains=search)
            )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search_query'] = self.request.GET.get('search', '')
        return context
#

class DepartmentDetailView(DetailView):
    model = Department
    template_name = 'store_management/department_detail.html'  # Your template path
    context_object_name = 'department'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Add store to context from the department's foreign key
        context['store'] = self.object.store
        # Add employees in this department
        # Removed employees context as per request
        # context['employees'] = self.object.employees.all()
        return context


def store_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = Store.objects.all()

    # Apply filters if any
    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(name__icontains=search) | Q(city__icontains=search) | Q(region__icontains=search)
        )

    region = request.GET.get('region')
    if region:
        queryset = queryset.filter(region=region)

    if format_type == 'excel':
        # Create Excel response
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="stores.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Stores'

        # Header
        headers = ['Store Name', 'Address', 'City', 'Region', 'Postal Code', 'Phone', 'Opening Date', 'Manager']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        # Data
        for row_num, store in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=store.name)
            ws.cell(row=row_num, column=2, value=store.address)
            ws.cell(row=row_num, column=3, value=store.city)
            ws.cell(row=row_num, column=4, value=store.region)
            ws.cell(row=row_num, column=5, value=store.postal_code or '')
            ws.cell(row=row_num, column=6, value=store.phone)
            ws.cell(row=row_num, column=7, value=store.opening_date.strftime('%Y-%m-%d'))
            ws.cell(row=row_num, column=8, value=store.manager.get_full_name() if store.manager else 'Unassigned')

        wb.save(response)
        return response

    else:  # Default to CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="stores.csv"'

        writer = csv.writer(response)
        writer.writerow(['Store Name', 'Address', 'City', 'Region', 'Postal Code', 'Phone', 'Opening Date', 'Manager'])

        for store in queryset:
            writer.writerow([
                store.name,
                store.address,
                store.city,
                store.region,
                store.postal_code or '',
                store.phone,
                store.opening_date.strftime('%Y-%m-%d'),
                store.manager.get_full_name() if store.manager else 'Unassigned',
            ])

        return response


def department_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = Department.objects.select_related('store').all()

    # Apply filters if any
    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(department_name__icontains=search) | Q(description__icontains=search) | Q(store__name__icontains=search)
        )

    store_id = request.GET.get('store')
    if store_id:
        queryset = queryset.filter(store_id=store_id)

    if format_type == 'excel':
        # Create Excel response
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="departments.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Departments'

        # Header
        headers = ['Department Name', 'Description', 'Store', 'Created At']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        # Data
        for row_num, department in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=department.department_name)
            ws.cell(row=row_num, column=2, value=department.description or '')
            ws.cell(row=row_num, column=3, value=department.store.name)
            ws.cell(row=row_num, column=4, value=department.created_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # Default to CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="departments.csv"'

        writer = csv.writer(response)
        writer.writerow(['Department Name', 'Description', 'Store', 'Created At'])

        for department in queryset:
            writer.writerow([
                department.department_name,
                department.description or '',
                department.store.name,
                department.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])

        return response


def generate_store_pdf(request, pk):
    """Generate PDF with store details and barcode."""
    store = get_object_or_404(Store, pk=pk)

    # Create PDF response
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="store_{store.store_number}.pdf"'

    # Create PDF
    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter

    # Title
    p.setFont("Helvetica-Bold", 16)
    p.drawString(100, height - 50, "Store Information")

    # Store details
    p.setFont("Helvetica", 12)
    y_position = height - 80

    p.drawString(100, y_position, f"Store Number: {store.store_number}")
    y_position -= 20
    p.drawString(100, y_position, f"Store Name: {store.name}")
    y_position -= 20
    p.drawString(100, y_position, f"Address: {store.address}")
    y_position -= 20
    p.drawString(100, y_position, f"City: {store.city}")
    y_position -= 20
    p.drawString(100, y_position, f"Region: {store.region}")
    y_position -= 20
    p.drawString(100, y_position, f"Phone: {store.phone}")
    y_position -= 20
    p.drawString(100, y_position, f"Opening Date: {store.opening_date.strftime('%Y-%m-%d')}")
    y_position -= 20
    p.drawString(100, y_position, f"Manager: {store.manager.get_full_name() if store.manager else 'Unassigned'}")

    # Generate and add barcode
    try:
        barcode_buffer = store.generate_barcode_image()
        barcode_image = ImageReader(barcode_buffer)
        p.drawImage(barcode_image, 100, y_position - 150, width=200, height=100)
        p.drawString(100, y_position - 170, f"Barcode: {store.store_number}")
    except Exception as e:
        p.drawString(100, y_position - 150, f"Barcode generation failed: {str(e)}")

    # Save PDF
    p.showPage()
    p.save()

    # Get PDF data from buffer
    pdf = buffer.getvalue()
    buffer.close()
    response.write(pdf)

    return response
