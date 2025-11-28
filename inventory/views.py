from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.auth.models import User
from django.contrib.messages.views import SuccessMessageMixin
from django.contrib import messages
from django.urls import reverse_lazy
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.db.models import Q, F, Sum, Count
from django.http import HttpResponse
from django.utils import timezone
from django.db import transaction
import csv
try:
    from openpyxl import Workbook
except Exception:
    Workbook = None

from .models import (
    Brand, Category, Product, InventoryRecord, StockMovement,
    InventoryAlert, ExpirationConfig, ReplenishmentConfig, StockTransfer
)
from .forms import BrandForm, CategoryForm, ProductForm, InventoryRecordForm, StockMovementForm, StockTransferForm
from .replenishment_utils import process_replenishment_batch
from store_management.models import Store


# Brand CRUD Views
class BrandListView(LoginRequiredMixin, ListView):
    model = Brand
    template_name = 'inventory/brand_list.html'
    context_object_name = 'brands'
    paginate_by = 25

    def get_queryset(self):
        queryset = super().get_queryset()
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) | Q(description__icontains=search)
            )
        return queryset.order_by('name')


class BrandCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = Brand
    form_class = BrandForm
    template_name = 'inventory/brand_form.html'
    success_url = reverse_lazy('inventory:brand-list')
    success_message = "Brand created successfully."


class BrandUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = Brand
    form_class = BrandForm
    template_name = 'inventory/brand_form.html'
    success_url = reverse_lazy('inventory:brand-list')
    success_message = "Brand updated successfully."


class BrandDeleteView(LoginRequiredMixin, DeleteView):
    model = Brand
    template_name = 'inventory/brand_confirm_delete.html'
    success_url = reverse_lazy('inventory:brand-list')

    def delete(self, request, *args, **kwargs):
        messages.success(request, "Brand deleted successfully.")
        return super().delete(request, *args, **kwargs)


@login_required
def brand_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = Brand.objects.all()

    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(name__icontains=search) | Q(description__icontains=search)
        )

    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="brands.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Brands'

        headers = ['Name', 'Description', 'Created At', 'Updated At']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        for row_num, brand in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=brand.name)
            ws.cell(row=row_num, column=2, value=brand.description or '')
            ws.cell(row=row_num, column=3, value=brand.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row_num, column=4, value=brand.updated_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="brands.csv"'

        writer = csv.writer(response)
        writer.writerow(['Name', 'Description', 'Created At', 'Updated At'])

        for brand in queryset:
            writer.writerow([
                brand.name,
                brand.description or '',
                brand.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                brand.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])

        return response


# Category CRUD Views
class CategoryListView(LoginRequiredMixin, ListView):
    model = Category
    template_name = 'inventory/category_list.html'
    context_object_name = 'categories'
    paginate_by = 25

    def get_queryset(self):
        queryset = super().get_queryset()
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) | Q(description__icontains=search)
            )
        return queryset.order_by('name')


class CategoryCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = Category
    form_class = CategoryForm
    template_name = 'inventory/category_form.html'
    success_url = reverse_lazy('inventory:category-list')
    success_message = "Category created successfully."


class CategoryUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = Category
    form_class = CategoryForm
    template_name = 'inventory/category_form.html'
    success_url = reverse_lazy('inventory:category-list')
    success_message = "Category updated successfully."


class CategoryDeleteView(LoginRequiredMixin, DeleteView):
    model = Category
    template_name = 'inventory/category_confirm_delete.html'
    success_url = reverse_lazy('inventory:category-list')

    def delete(self, request, *args, **kwargs):
        messages.success(request, "Category deleted successfully.")
        return super().delete(request, *args, **kwargs)


@login_required
def category_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = Category.objects.all()

    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(name__icontains=search) | Q(description__icontains=search)
        )

    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="categories.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Categories'

        headers = ['Name', 'Description', 'Created At', 'Updated At']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        for row_num, category in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=category.name)
            ws.cell(row=row_num, column=2, value=category.description or '')
            ws.cell(row=row_num, column=3, value=category.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row_num, column=4, value=category.updated_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="categories.csv"'

        writer = csv.writer(response)
        writer.writerow(['Name', 'Description', 'Created At', 'Updated At'])

        for category in queryset:
            writer.writerow([
                category.name,
                category.description or '',
                category.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                category.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])

        return response


# Product CRUD Views
class ProductListView(LoginRequiredMixin, ListView):
    model = Product
    template_name = 'inventory/product_list.html'
    context_object_name = 'products'
    paginate_by = 25

    def get_queryset(self):
        queryset = super().get_queryset().select_related('category', 'brand', 'default_supplier')
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) | Q(sku__icontains=search) | Q(description__icontains=search)
            )

        category = self.request.GET.get('category')
        if category:
            queryset = queryset.filter(category_id=category)

        brand = self.request.GET.get('brand')
        if brand:
            queryset = queryset.filter(brand_id=brand)

        is_active = self.request.GET.get('is_active')
        if is_active is not None:
            queryset = queryset.filter(is_active=is_active == 'true')

        return queryset.order_by('name')


class ProductDetailView(LoginRequiredMixin, DetailView):
    model = Product
    template_name = 'inventory/product_detail.html'
    context_object_name = 'product'


class ProductCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = Product
    form_class = ProductForm
    template_name = 'inventory/product_form.html'
    success_url = reverse_lazy('inventory:product-list')
    success_message = "Product created successfully."


class ProductUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = Product
    form_class = ProductForm
    template_name = 'inventory/product_form.html'
    success_url = reverse_lazy('inventory:product-list')
    success_message = "Product updated successfully."


class ProductDeleteView(LoginRequiredMixin, DeleteView):
    model = Product
    template_name = 'inventory/product_confirm_delete.html'
    success_url = reverse_lazy('inventory:product-list')

    def delete(self, request, *args, **kwargs):
        messages.success(request, "Product deleted successfully.")
        return super().delete(request, *args, **kwargs)


@login_required
def product_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = Product.objects.select_related('category', 'brand', 'default_supplier').all()

    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(name__icontains=search) | Q(sku__icontains=search) | Q(description__icontains=search)
        )

    category = request.GET.get('category')
    if category:
        queryset = queryset.filter(category_id=category)

    brand = request.GET.get('brand')
    if brand:
        queryset = queryset.filter(brand_id=brand)

    is_active = request.GET.get('is_active')
    if is_active is not None:
        queryset = queryset.filter(is_active=is_active == 'true')

    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="products.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Products'

        headers = ['Name', 'SKU', 'Category', 'Brand', 'Unit Price', 'Reorder Level', 'Criticality', 'Active', 'Barcode', 'Created At']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        for row_num, product in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=product.name)
            ws.cell(row=row_num, column=2, value=product.sku or '')
            ws.cell(row=row_num, column=3, value=product.category.name)
            ws.cell(row=row_num, column=4, value=product.brand.name)
            ws.cell(row=row_num, column=5, value=float(product.unit_price))
            ws.cell(row=row_num, column=6, value=product.reorder_level)
            ws.cell(row=row_num, column=7, value=product.get_criticality_display())
            ws.cell(row=row_num, column=8, value='Yes' if product.is_active else 'No')
            ws.cell(row=row_num, column=9, value=product.barcode or '')
            ws.cell(row=row_num, column=10, value=product.created_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="products.csv"'

        writer = csv.writer(response)
        writer.writerow(['Name', 'SKU', 'Category', 'Brand', 'Unit Price', 'Reorder Level', 'Criticality', 'Active', 'Barcode', 'Created At'])

        for product in queryset:
            writer.writerow([
                product.name,
                product.sku or '',
                product.category.name,
                product.brand.name,
                product.unit_price,
                product.reorder_level,
                product.get_criticality_display(),
                'Yes' if product.is_active else 'No',
                product.barcode or '',
                product.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])

        return response


# InventoryRecord CRUD Views
class InventoryRecordListView(LoginRequiredMixin, ListView):
    model = InventoryRecord
    template_name = 'inventory/inventoryrecord_list.html'
    context_object_name = 'inventory_records'
    paginate_by = 25

    def get_queryset(self):
        queryset = super().get_queryset().select_related('product', 'store', 'supplier')
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(store__name__icontains=search)
            )

        store = self.request.GET.get('store')
        if store:
            queryset = queryset.filter(store_id=store)

        product = self.request.GET.get('product')
        if product:
            queryset = queryset.filter(product_id=product)

        return queryset.order_by('-updated_at')


class InventoryRecordCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = InventoryRecord
    form_class = InventoryRecordForm
    template_name = 'inventory/inventoryrecord_form.html'
    success_url = reverse_lazy('inventory:inventoryrecord-list')
    success_message = "Inventory record created successfully."


class InventoryRecordUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = InventoryRecord
    form_class = InventoryRecordForm
    template_name = 'inventory/inventoryrecord_form.html'
    success_url = reverse_lazy('inventory:inventoryrecord-list')
    success_message = "Inventory record updated successfully."


class InventoryRecordDeleteView(LoginRequiredMixin, DeleteView):
    model = InventoryRecord
    template_name = 'inventory/inventoryrecord_confirm_delete.html'
    success_url = reverse_lazy('inventory:inventoryrecord-list')

    def delete(self, request, *args, **kwargs):
        messages.success(request, "Inventory record deleted successfully.")
        return super().delete(request, *args, **kwargs)


@login_required
def inventoryrecord_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = InventoryRecord.objects.select_related('product', 'store', 'supplier').all()

    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(store__name__icontains=search)
        )

    store = request.GET.get('store')
    if store:
        queryset = queryset.filter(store_id=store)

    product = request.GET.get('product')
    if product:
        queryset = queryset.filter(product_id=product)

    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="inventory_records.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Inventory Records'

        headers = ['Product', 'SKU', 'Store', 'Quantity', 'Location', 'Batch Number', 'Expiration Date', 'Cost Price', 'Supplier', 'Last Updated']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        for row_num, record in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=record.product.name)
            ws.cell(row=row_num, column=2, value=record.product.sku or '')
            ws.cell(row=row_num, column=3, value=record.store.name)
            ws.cell(row=row_num, column=4, value=record.quantity)
            ws.cell(row=row_num, column=5, value=record.location or '')
            ws.cell(row=row_num, column=6, value=record.batch_number or '')
            ws.cell(row=row_num, column=7, value=record.expiration_date.strftime('%Y-%m-%d') if record.expiration_date else '')
            ws.cell(row=row_num, column=8, value=float(record.cost_price) if record.cost_price else '')
            ws.cell(row=row_num, column=9, value=record.supplier.name if record.supplier else '')
            ws.cell(row=row_num, column=10, value=record.updated_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="inventory_records.csv"'

        writer = csv.writer(response)
        writer.writerow(['Product', 'SKU', 'Store', 'Quantity', 'Location', 'Batch Number', 'Expiration Date', 'Cost Price', 'Supplier', 'Last Updated'])

        for record in queryset:
            writer.writerow([
                record.product.name,
                record.product.sku or '',
                record.store.name,
                record.quantity,
                record.location or '',
                record.batch_number or '',
                record.expiration_date.strftime('%Y-%m-%d') if record.expiration_date else '',
                record.cost_price if record.cost_price else '',
                record.supplier.name if record.supplier else '',
                record.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])


# StockTransfer CRUD Views
class StockTransferListView(LoginRequiredMixin, ListView):
    model = StockTransfer
    template_name = 'inventory/stocktransfer_list.html'
    context_object_name = 'stocktransfers'
    paginate_by = 25

    def get_queryset(self):
        queryset = super().get_queryset().select_related('product', 'from_store', 'to_store', 'created_by')
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(from_store__name__icontains=search) | Q(to_store__name__icontains=search)
            )

        from_store = self.request.GET.get('from_store')
        if from_store:
            queryset = queryset.filter(from_store_id=from_store)

        to_store = self.request.GET.get('to_store')
        if to_store:
            queryset = queryset.filter(to_store_id=to_store)

        created_by = self.request.GET.get('created_by')
        if created_by:
            queryset = queryset.filter(created_by_id=created_by)

        return queryset.order_by('-created_at')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['total_transfers'] = self.get_queryset().count()
        context['total_quantity'] = self.get_queryset().aggregate(Sum('quantity'))['quantity__sum'] or 0
        context['recent_transfers'] = self.get_queryset().filter(created_at__gte=timezone.now() - timezone.timedelta(days=7)).count()
        context['stores'] = Store.objects.all()
        context['users'] = User.objects.all()
        return context


class StockTransferCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = StockTransfer
    form_class = StockTransferForm
    template_name = 'inventory/stocktransfer_form.html'
    success_url = reverse_lazy('inventory:stocktransfer-list')
    success_message = "Stock transfer created successfully."

    def form_valid(self, form):
        form.instance.created_by = self.request.user.employee_profile
        return super().form_valid(form)


class StockTransferUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = StockTransfer
    form_class = StockTransferForm
    template_name = 'inventory/stocktransfer_form.html'
    success_url = reverse_lazy('inventory:stocktransfer-list')
    success_message = "Stock transfer updated successfully."


class StockTransferDeleteView(LoginRequiredMixin, DeleteView):
    model = StockTransfer
    template_name = 'inventory/stocktransfer_confirm_delete.html'
    success_url = reverse_lazy('inventory:stocktransfer-list')

    def delete(self, request, *args, **kwargs):
        messages.success(request, "Stock transfer deleted successfully.")
        return super().delete(request, *args, **kwargs)


class StockTransferDetailView(LoginRequiredMixin, DetailView):
    model = StockTransfer
    template_name = 'inventory/stocktransfer_detail.html'
    context_object_name = 'stocktransfer'


@login_required
def stocktransfer_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = StockTransfer.objects.select_related('product', 'from_store', 'to_store', 'created_by').all()

    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(from_store__name__icontains=search) | Q(to_store__name__icontains=search)
        )

    from_store = request.GET.get('from_store')
    if from_store:
        queryset = queryset.filter(from_store_id=from_store)

    to_store = request.GET.get('to_store')
    if to_store:
        queryset = queryset.filter(to_store_id=to_store)

    product = request.GET.get('product')
    if product:
        queryset = queryset.filter(product_id=product)

    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="stock_transfers.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Stock Transfers'

        headers = ['Product', 'SKU', 'From Store', 'To Store', 'Quantity', 'Notes', 'Created By', 'Created At']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        for row_num, transfer in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=transfer.product.name)
            ws.cell(row=row_num, column=2, value=transfer.product.sku or '')
            ws.cell(row=row_num, column=3, value=transfer.from_store.name)
            ws.cell(row=row_num, column=4, value=transfer.to_store.name)
            ws.cell(row=row_num, column=5, value=transfer.quantity)
            ws.cell(row=row_num, column=6, value=transfer.notes or '')
            ws.cell(row=row_num, column=7, value=str(transfer.created_by) if transfer.created_by else '')
            ws.cell(row=row_num, column=8, value=transfer.created_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="stock_transfers.csv"'

        writer = csv.writer(response)
        writer.writerow(['Product', 'SKU', 'From Store', 'To Store', 'Quantity', 'Notes', 'Created By', 'Created At'])

        for transfer in queryset:
            writer.writerow([
                transfer.product.name,
                transfer.product.sku or '',
                transfer.from_store.name,
                transfer.to_store.name,
                transfer.quantity,
                transfer.notes or '',
                str(transfer.created_by) if transfer.created_by else '',
                transfer.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])

        return response


# StockMovement CRUD Views
class StockMovementListView(LoginRequiredMixin, ListView):
    model = StockMovement
    template_name = 'inventory/stockmovement_list.html'
    context_object_name = 'stock_movements'
    paginate_by = 25

    def get_queryset(self):
        queryset = super().get_queryset().select_related('product', 'store', 'created_by')
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(store__name__icontains=search) | Q(reference__icontains=search)
            )

        store = self.request.GET.get('store')
        if store:
            queryset = queryset.filter(store_id=store)

        product = self.request.GET.get('product')
        if product:
            queryset = queryset.filter(product_id=product)

        movement_type = self.request.GET.get('movement_type')
        if movement_type:
            queryset = queryset.filter(movement_type=movement_type)

        return queryset.order_by('-created_at')


class StockMovementCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = StockMovement
    form_class = StockMovementForm
    template_name = 'inventory/stockmovement_form.html'
    success_url = reverse_lazy('inventory:stockmovement-list')
    success_message = "Stock movement created successfully."

    def form_valid(self, form):
        form.instance.created_by = self.request.user.employee_profile
        return super().form_valid(form)


class StockMovementUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = StockMovement
    form_class = StockMovementForm
    template_name = 'inventory/stockmovement_form.html'
    success_url = reverse_lazy('inventory:stockmovement-list')
    success_message = "Stock movement updated successfully."


class StockMovementDeleteView(LoginRequiredMixin, DeleteView):
    model = StockMovement
    template_name = 'inventory/stockmovement_confirm_delete.html'
    success_url = reverse_lazy('inventory:stockmovement-list')

    def delete(self, request, *args, **kwargs):
        messages.success(request, "Stock movement deleted successfully.")
        return super().delete(request, *args, **kwargs)


@login_required
def stockmovement_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = StockMovement.objects.select_related('product', 'store', 'created_by').all()

    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(store__name__icontains=search) | Q(reference__icontains=search)
        )

    store = request.GET.get('store')
    if store:
        queryset = queryset.filter(store_id=store)

    product = self.request.GET.get('product')
    if product:
        queryset = queryset.filter(product_id=product)

    movement_type = request.GET.get('movement_type')
    if movement_type:
        queryset = queryset.filter(movement_type=movement_type)

    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="stock_movements.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Stock Movements'

        headers = ['Product', 'SKU', 'Store', 'Quantity', 'Movement Type', 'Reason', 'Reference', 'Created By', 'Created At']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        for row_num, movement in enumerate(queryset, 2):
            ws.cell(row=row_num, column=1, value=movement.product.name)
            ws.cell(row=row_num, column=2, value=movement.product.sku or '')
            ws.cell(row=row_num, column=3, value=movement.store.name)
            ws.cell(row=row_num, column=4, value=movement.quantity)
            ws.cell(row=row_num, column=5, value=movement.get_movement_type_display())
            ws.cell(row=row_num, column=6, value=movement.get_reason_display() if movement.reason else '')
            ws.cell(row=row_num, column=7, value=movement.reference or '')
            ws.cell(row=row_num, column=8, value=str(movement.created_by) if movement.created_by else '')
            ws.cell(row=row_num, column=9, value=movement.created_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="stock_movements.csv"'

        writer = csv.writer(response)
        writer.writerow(['Product', 'SKU', 'Store', 'Quantity', 'Movement Type', 'Reason', 'Reference', 'Created By', 'Created At'])

        for movement in queryset:
            writer.writerow([
                movement.product.name,
                movement.product.sku or '',
                movement.store.name,
                movement.quantity,
                movement.get_movement_type_display(),
                movement.get_reason_display() if movement.reason else '',
                movement.reference or '',
                str(movement.created_by) if movement.created_by else '',
                movement.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])

        return response


# Dashboard Views
class InventoryDashboardView(LoginRequiredMixin, ListView):
    model = InventoryRecord
    template_name = 'inventory/inventory_dashboard.html'
    context_object_name = 'inventory_records'

    def get_queryset(self):
        return InventoryRecord.objects.select_related('product', 'store').filter(
            quantity__lte=F('product__reorder_level')
        ).order_by('quantity')[:20]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # Low stock alerts
        context['low_stock_count'] = InventoryRecord.objects.filter(
            quantity__lte=F('product__reorder_level'), quantity__gt=0
        ).count()

        # Out of stock alerts
        context['out_of_stock_count'] = InventoryRecord.objects.filter(quantity=0).count()

        # Total inventory value
        context['total_inventory_value'] = InventoryRecord.objects.aggregate(
            total=Sum(F('quantity') * F('product__unit_price'))
        )['total'] or 0

        # Recent stock movements
        context['recent_movements'] = StockMovement.objects.select_related(
            'product', 'store'
        ).order_by('-created_at')[:10]

        return context


class StockManagementView(LoginRequiredMixin, ListView):
    model = InventoryRecord
    template_name = 'inventory/stock_management.html'
    context_object_name = 'inventory_records'
    paginate_by = 50

    def get_queryset(self):
        queryset = InventoryRecord.objects.select_related('product', 'store', 'supplier').all()

        # Apply filters
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(store__name__icontains=search)
            )

        stores = self.request.GET.getlist('store')
        if stores:
            queryset = queryset.filter(store_id__in=stores)

        categories = self.request.GET.getlist('category')
        if categories:
            queryset = queryset.filter(product__category_id__in=categories)

        criticality = self.request.GET.getlist('criticality')
        if criticality:
            queryset = queryset.filter(product__criticality__in=criticality)

        product = self.request.GET.get('product')
        if product:
            queryset = queryset.filter(product_id=product)

        stock_status = self.request.GET.get('stock_status')
        if stock_status == 'low_stock':
            queryset = queryset.filter(quantity__lte=F('product__reorder_level'), quantity__gt=0)
        elif stock_status == 'out_of_stock':
            queryset = queryset.filter(quantity=0)
        elif stock_status == 'overstock':
            queryset = queryset.filter(quantity__gt=F('product__reorder_level') * 2)

        return queryset.order_by('product__name')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # Add stock status to each record
        for record in context['inventory_records']:
            if record.quantity == 0:
                record.stock_status = 'out_of_stock'
            elif record.quantity <= record.product.reorder_level:
                record.stock_status = 'low_stock'
            elif record.quantity > record.product.reorder_level * 2:
                record.stock_status = 'overstock'
            else:
                record.stock_status = 'in_stock'

        return context


class ExpirationDashboardView(LoginRequiredMixin, ListView):
    model = InventoryRecord
    template_name = 'inventory/expiration_dashboard.html'
    context_object_name = 'expiring_items'

    def get_queryset(self):
        today = timezone.now().date()
        config = ExpirationConfig.get_default_config()

        # Get items expiring within the critical threshold
        critical_threshold = today + timezone.timedelta(days=config.get_threshold_days('critical'))
        queryset = InventoryRecord.objects.select_related('product', 'store').filter(
            expiration_date__lte=critical_threshold,
            expiration_date__gte=today,
            quantity__gt=0
        ).order_by('expiration_date')

        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = timezone.now().date()
        config = ExpirationConfig.get_default_config()

        # Expired items
        context['expired_items'] = InventoryRecord.objects.select_related('product', 'store').filter(
            expiration_date__lt=today,
            quantity__gt=0
        ).order_by('expiration_date')

        # Items expiring soon (warning threshold)
        warning_threshold = today + timezone.timedelta(days=config.get_threshold_days('warning'))
        context['warning_items'] = InventoryRecord.objects.select_related('product', 'store').filter(
            expiration_date__lte=warning_threshold,
            expiration_date__gt=today + timezone.timedelta(days=config.get_threshold_days('critical')),
            quantity__gt=0
        ).order_by('expiration_date')

        # Items expiring in future (info threshold)
        info_threshold = today + timezone.timedelta(days=config.get_threshold_days('info'))
        context['info_items'] = InventoryRecord.objects.select_related('product', 'store').filter(
            expiration_date__lte=info_threshold,
            expiration_date__gt=warning_threshold,
            quantity__gt=0
        ).order_by('expiration_date')

        context['config'] = config
        return context


@login_required
def process_replenishment(request):
    if request.method == 'POST':
        try:
            with transaction.atomic():
                batch = process_replenishment_batch()
                messages.success(request, f"Replenishment processed successfully. Batch ID: {batch.batch_id}")
                return redirect('inventory:stock-management')
        except Exception as e:
            messages.error(request, f"Error processing replenishment: {str(e)}")
            return redirect('inventory:stock-management')

    return render(request, 'inventory/process_replenishment.html')


@login_required
@permission_required('inventory.view_inventoryrecord', raise_exception=True)
def centralized_stock_export(request):
    format_type = request.GET.get('format', 'csv').lower()
    queryset = InventoryRecord.objects.select_related('product', 'store', 'supplier').all()

    # Apply filters
    search = request.GET.get('search')
    if search:
        queryset = queryset.filter(
            Q(product__name__icontains=search) | Q(product__sku__icontains=search) | Q(store__name__icontains=search)
        )

    stores = request.GET.getlist('store')
    if stores:
        queryset = queryset.filter(store_id__in=stores)

    categories = request.GET.getlist('category')
    if categories:
        queryset = queryset.filter(product__category_id__in=categories)

    criticality = request.GET.getlist('criticality')
    if criticality:
        queryset = queryset.filter(product__criticality__in=criticality)

    stock_status = request.GET.get('stock_status')
    if stock_status == 'low_stock':
        queryset = queryset.filter(quantity__lte=F('product__reorder_level'), quantity__gt=0)
    elif stock_status == 'out_of_stock':
        queryset = queryset.filter(quantity=0)
    elif stock_status == 'overstock':
        queryset = queryset.filter(quantity__gt=F('product__reorder_level') * 2)

    if format_type == 'excel':
        # Create Excel response
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="centralized_stock_report.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Centralized Stock Report'

        # Header
        headers = ['Product', 'SKU', 'Category', 'Store', 'Quantity', 'Reorder Level', 'Stock Status', 'Criticality', 'Unit Price', 'Total Value', 'Supplier', 'Last Updated']
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)

        # Data
        for row_num, record in enumerate(queryset, 2):
            # Determine stock status
            if record.quantity == 0:
                status = 'Out of Stock'
            elif record.quantity <= record.product.reorder_level:
                status = 'Low Stock'
            elif record.quantity > record.product.reorder_level * 2:
                status = 'Overstock'
            else:
                status = 'In Stock'

            ws.cell(row=row_num, column=1, value=record.product.name)
            ws.cell(row=row_num, column=2, value=record.product.sku or '')
            ws.cell(row=row_num, column=3, value=record.product.category.name)
            ws.cell(row=row_num, column=4, value=record.store.name)
            ws.cell(row=row_num, column=5, value=record.quantity)
            ws.cell(row=row_num, column=6, value=record.product.reorder_level)
            ws.cell(row=row_num, column=7, value=status)
            ws.cell(row=row_num, column=8, value=record.product.get_criticality_display())
            ws.cell(row=row_num, column=9, value=float(record.product.unit_price))
            ws.cell(row=row_num, column=10, value=float(record.quantity * record.product.unit_price))
            ws.cell(row=row_num, column=11, value=record.supplier.name if record.supplier else '')
            ws.cell(row=row_num, column=12, value=record.updated_at.strftime('%Y-%m-%d %H:%M:%S'))

        wb.save(response)
        return response

    else:  # Default to CSV
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="centralized_stock_report.csv"'

        writer = csv.writer(response)
        writer.writerow(['Product', 'SKU', 'Category', 'Store', 'Quantity', 'Reorder Level', 'Stock Status', 'Criticality', 'Unit Price', 'Total Value', 'Supplier', 'Last Updated'])

        for record in queryset:
            # Determine stock status
            if record.quantity == 0:
                status = 'Out of Stock'
            elif record.quantity <= record.product.reorder_level:
                status = 'Low Stock'
            elif record.quantity > record.product.reorder_level * 2:
                status = 'Overstock'
            else:
                status = 'In Stock'

            writer.writerow([
                record.product.name,
                record.product.sku or '',
                record.product.category.name,
                record.store.name,
                record.quantity,
                record.product.reorder_level,
                status,
                record.product.get_criticality_display(),
                record.product.unit_price,
                record.quantity * record.product.unit_price,
                record.supplier.name if record.supplier else '',
                record.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
            ])
